"""
Evidence file handling: text extraction from documents.
"""
from io import BytesIO

import PyPDF2
from docx import Document

DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".csv", ".log", ".json"}
IMAGE_TYPES = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png", ".webp": "image/webp"}


def extension(filename: str) -> str:
    name = (filename or "").lower()
    return name[name.rfind("."):] if "." in name else ""


def extract_document_text(filename: str, data: bytes) -> str:
    ext = extension(filename)
    if ext == ".pdf":
        reader = PyPDF2.PdfReader(BytesIO(data))
        pages = []
        for i, page in enumerate(reader.pages, 1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Page {i}]\n{text}")
        return "\n\n".join(pages)
    if ext == ".docx":
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""
